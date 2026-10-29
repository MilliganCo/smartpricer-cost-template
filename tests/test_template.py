import hashlib
import importlib.util
import contextlib
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

PROJECT = Path(__file__).parents[1]
spec = importlib.util.spec_from_file_location("cost_template", PROJECT / "sebes.py")
sebes = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sebes)


class TemplateTests(unittest.TestCase):
    def setUp(self):
        work = PROJECT / "work"
        work.mkdir(exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(dir=work)
        self.folder = Path(self.temporary.name)
        self.old_tempdir, tempfile.tempdir = tempfile.tempdir, str(self.folder)
        self.old_cwd = Path.cwd()
        os.chdir(self.folder)
        self.paths = {
            "ozon": self.folder / "costs.xlsx", "articles": self.folder / "articles.xlsx",
            "template": self.folder / "template.xlsx", "output": self.folder / "output.xlsx",
        }
        self.globals = patch.multiple(sebes, OZON_PRICES=self.paths["ozon"], ARTICUL_PROFIT=self.paths["articles"], SMARTPRICER_TEMPLATE=self.paths["template"], RESULT_FILE=self.paths["output"])
        self.globals.start()

    def tearDown(self):
        self.globals.stop()
        os.chdir(self.old_cwd)
        tempfile.tempdir = self.old_tempdir
        self.temporary.cleanup()

    def fixtures(self, duplicate=False, missing_cost=False, text_row_keys=False):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Costs"
        sheet.append(["НОМЕР СТРОКИ", "СЕБЕСТОИМОСТЬ"])
        sheet.append([1, 12.5])
        sheet.append([2, None if missing_cost else 30])
        if duplicate:
            sheet.append([1, 99])
        workbook.save(self.paths["ozon"])
        workbook.close()
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["НОМЕР СТРОКИ", "Артикул Ozon"])
        sheet.append(["01" if text_row_keys else 1, "0012345678905"])
        sheet.append(["2.0" if text_row_keys else 2, "12345678905"])
        workbook.save(self.paths["articles"])
        workbook.close()
        workbook = Workbook()
        sheet = workbook.active
        for _ in range(5):
            sheet.append(["synthetic instructions"])
        sheet.append(["Артикул товара", "Закупочная цена"])
        workbook.save(self.paths["template"])
        workbook.close()

    def run_program(self, *extra):
        arguments = ["sebes.py", "--ozon-prices", str(self.paths["ozon"]), "--articles", str(self.paths["articles"]), "--template", str(self.paths["template"]), "--out", str(self.paths["output"])]
        arguments.extend(extra)
        with patch("sys.argv", arguments):
            return sebes.main()

    def test_only_temp_output_preserves_text_articles_costs_and_inputs(self):
        self.fixtures()
        before = {key: hashlib.sha256(path.read_bytes()).hexdigest() for key, path in self.paths.items() if key != "output"}
        self.run_program()
        result = load_workbook(self.paths["output"], read_only=True, data_only=True)
        rows = list(result.active.values)
        result.close()
        self.assertEqual(rows[0], ("Артикул товара", "Закупочная цена"))
        self.assertEqual(rows[1:], [("0012345678905", 12.5), ("12345678905", 30)])
        self.assertEqual(before, {key: hashlib.sha256(path.read_bytes()).hexdigest() for key, path in self.paths.items() if key != "output"})

    def test_ambiguous_duplicate_rows_are_rejected_before_output(self):
        self.fixtures(duplicate=True)
        with self.assertRaises(ValueError):
            self.run_program()
        self.assertFalse(self.paths["output"].exists())

    def test_missing_cost_stays_blank_instead_of_zero(self):
        self.fixtures(missing_cost=True)
        self.run_program()
        result = load_workbook(self.paths["output"], read_only=True, data_only=True)
        rows = list(result.active.values)
        result.close()
        self.assertIsNone(rows[2][1])

    def test_output_cannot_replace_an_input(self):
        self.fixtures()
        before = self.paths["ozon"].read_bytes()
        self.paths["output"] = self.paths["ozon"]
        with self.assertRaises(ValueError):
            self.run_program("--overwrite")
        self.assertEqual(self.paths["ozon"].read_bytes(), before)

    def test_existing_output_requires_explicit_overwrite(self):
        self.fixtures()
        self.paths["output"].write_bytes(b"synthetic existing output")
        with self.assertRaises(FileExistsError):
            self.run_program()
        self.assertEqual(self.paths["output"].read_bytes(), b"synthetic existing output")
        self.assertEqual(self.run_program("--overwrite"), 0)

    def test_explicit_config_relative_paths_and_empty_example(self):
        self.fixtures()
        config = self.folder / "private.json"
        config.write_text(json.dumps({k: "" for k in ("ozon_prices", "articles", "template", "out")}), encoding="utf-8")
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as result:
            sebes.main(["--config", str(config)])
        self.assertEqual(result.exception.code, 2)
        self.assertFalse(self.paths["output"].exists())
        config.write_text(json.dumps({"ozon_prices": self.paths["ozon"].name, "articles": self.paths["articles"].name, "template": self.paths["template"].name, "out": self.paths["output"].name}), encoding="utf-8")
        self.assertEqual(sebes.main(["--config", str(config)]), 0)
        self.assertTrue(self.paths["output"].is_file())

    def test_integer_text_row_keys_match_without_losing_article_text(self):
        self.fixtures(text_row_keys=True)
        self.run_program()
        result = load_workbook(self.paths["output"], read_only=True, data_only=True)
        rows = list(result.active.values)
        result.close()
        self.assertEqual([row[0] for row in rows[1:]], ["0012345678905", "12345678905"])


if __name__ == "__main__":
    unittest.main()
